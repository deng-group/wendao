from __future__ import annotations

import contextlib
import io
import json
import os
import tempfile
import threading
import unittest
import zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from unittest import mock

from wendao import workspace as workspace_module
from wendao.cli import main
from wendao.pack import PackError, open_pack, unpack
from wendao.rag.providers import OpenAICompatibleProvider
from wendao.web.ai import AiPolicy, AiUnavailable
from wendao.web.explorer import create_app as create_explorer
from wendao.web.widget import create_app as create_widget


class StudentKeyServer(BaseHTTPRequestHandler):
    """OpenAI-compatible endpoint that records which key it was called with."""

    keys: list[str] = []

    def do_POST(self):  # noqa: N802
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        StudentKeyServer.keys.append(self.headers.get("Authorization", ""))
        self.send_response(200)
        self.end_headers()
        if body.get("stream"):
            self.wfile.write(b'data: {"choices": [{"delta": {"content": "Kinetic energy is energy of motion."}}]}\n\ndata: [DONE]\n\n')
        else:
            self.wfile.write(b'{"choices": [{"message": {"content": "Kinetic energy is energy of motion."}}]}')

    def log_message(self, *args):
        pass


def run_cli(*argv: str) -> str:
    output = io.StringIO()
    with contextlib.redirect_stdout(output):
        main(list(argv))
    return output.getvalue()


def make_workspace(root: Path, student_settings: str) -> None:
    run_cli("init", str(root))
    (root / "notes" / "week1").mkdir(parents=True)
    (root / "notes" / "week2").mkdir(parents=True)
    (root / "notes" / "week1" / "laws.md").write_text(
        "# Newton's Laws\n\nNewton's second law relates force, mass, and acceleration: F = ma.\n", encoding="utf-8"
    )
    (root / "notes" / "week2" / "energy.md").write_text(
        "# Energy\n\nKinetic energy is the energy of motion. Work done by a force changes kinetic energy.\n", encoding="utf-8"
    )
    (root / "notes" / "syllabus.md").write_text("# Syllabus\n- Year: AY2026/2027 Semester 1\n\nThe final exam is in week 13.\n",
                                                encoding="utf-8")
    (root / "concepts.json").write_text(json.dumps({"version": 1, "concepts": [
        {"id": "kinetic-energy", "label": "Kinetic Energy", "category": "energy", "aliases": ["kinetic energy"]},
        {"id": "force", "label": "Force", "category": "mechanics", "aliases": ["force"]},
    ]}), encoding="utf-8")
    config = root / "wendao.toml"
    text = config.read_text(encoding="utf-8").replace("[search]", '[search]\nembedding_model = "tfidf"')
    start = text.index("[student]")
    end = text.index("[model]")
    config.write_text(text[:start] + "[student]\n" + student_settings + "\n\n" + text[end:], encoding="utf-8")


def stream_events(client, payload: dict) -> list[dict]:
    response = client.post("/api/answer/stream", json=payload)
    return [json.loads(line) for line in response.get_data(as_text=True).splitlines() if line.strip()]


class AiPolicyTest(unittest.TestCase):
    def setUp(self):
        self.env = mock.patch.dict(os.environ, {"OPENAI_API_KEY": "sk-teacher", "OPENAI_BASE_URL": "https://teacher.relay/v1"}, clear=True)
        self.env.start()

    def tearDown(self):
        self.env.stop()

    def test_teacher_mode_uses_teacher_key_with_daily_limit(self):
        policy = AiPolicy.from_settings("teacher", questions_per_day=2)
        for _ in range(2):
            how, provider = policy.choose({}, "student-a")
            self.assertEqual(how, "provider")
            self.assertTrue(provider.use_env)
        with self.assertRaisesRegex(AiUnavailable, "limit of 2"):
            policy.choose({}, "student-a")
        self.assertEqual(policy.choose({}, "student-b")[0], "provider")  # limits are per student

    def test_teacher_mode_refuses_student_keys(self):
        with self.assertRaisesRegex(AiUnavailable, "turned off"):
            AiPolicy.from_settings("teacher").choose({"ai": {"provider": "openai", "api_key": "sk-x"}}, "s")

    def test_student_key_never_mixes_with_teacher_settings(self):
        policy = AiPolicy.from_settings("either")
        _, provider = policy.choose({"ai": {"provider": "openai", "api_key": "sk-student"}}, "s")
        self.assertIsInstance(provider, OpenAICompatibleProvider)
        self.assertFalse(provider.use_env)
        request_obj = provider._request({"messages": [{"content": "s"}, {"content": "q"}]}, stream=False)
        self.assertEqual(request_obj.full_url, "https://api.openai.com/v1/chat/completions")
        self.assertEqual(request_obj.headers["Authorization"], "Bearer sk-student")

    def test_student_mode_requires_a_key(self):
        with self.assertRaisesRegex(AiUnavailable, "own AI key"):
            AiPolicy.from_settings("student").choose({}, "s")
        self.assertTrue(AiPolicy.from_settings("student").describe()["needs_student_key"])

    def test_course_app_ignores_keys_on_the_students_computer(self):
        policy = AiPolicy.from_settings("either", use_local_key=False)
        self.assertFalse(policy.describe()["teacher_ai"])
        with self.assertRaisesRegex(AiUnavailable, "not set up"):
            policy.choose({}, "s")
        self.assertEqual(AiPolicy.from_settings("teacher", server="https://course.example", use_local_key=False).choose({}, "s"),
                         ("forward", "https://course.example"))

    def test_bad_mode_is_rejected(self):
        with self.assertRaises(ValueError):
            AiPolicy.from_settings("sometimes")


class CourseAppTest(unittest.TestCase):
    """Teacher builds and packs; a student opens the file and asks with their own key."""

    @classmethod
    def setUpClass(cls):
        cls.server = HTTPServer(("127.0.0.1", 0), StudentKeyServer)
        threading.Thread(target=cls.server.serve_forever, daemon=True).start()
        cls.base_url = f"http://127.0.0.1:{cls.server.server_port}/v1"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "physics"
        self.env = mock.patch.dict(os.environ, {"WENDAO_CACHE": str(Path(self.tmp.name) / "cache")}, clear=True)
        self.env.start()

    def tearDown(self):
        self.env.stop()
        self.tmp.cleanup()

    def test_pack_and_open_with_student_key(self):
        make_workspace(self.root, 'ai = "either"')
        run_cli("build", "-w", str(self.root))
        output = run_cli("pack", "-w", str(self.root))
        pack_path = self.root / "build" / "physics.wendao"
        self.assertTrue(pack_path.is_file())
        self.assertIn("will need their own AI key", output)
        self.assertIn("wendao open physics.wendao", output)

        course = open_pack(pack_path)
        self.assertEqual(course.course_name, "Physics")
        self.assertFalse(course.uses_local_key)
        self.assertEqual(open_pack(pack_path).root, course.root, "opening twice reuses the unpacked folder")

        os.environ["OPENAI_API_KEY"] = "sk-on-the-laptop"  # must not be used as the course AI
        client = create_explorer(course).test_client()
        health = client.get("/api/health").get_json()
        self.assertTrue(health["ai"]["needs_student_key"])
        self.assertEqual(client.get("/api/graph").status_code, 200)
        self.assertIn("Wendao · Physics", client.get("/").get_data(as_text=True))

        events = stream_events(client, {"query": "What is kinetic energy?"})
        self.assertEqual(events[-1]["error"], "AiUnavailable")

        StudentKeyServer.keys.clear()
        events = stream_events(client, {"query": "What is kinetic energy?",
                                        "ai": {"provider": "openai", "base_url": self.base_url, "api_key": "sk-student"}})
        self.assertEqual(events[-1]["type"], "done")
        self.assertEqual(events[-1]["answer"], "Kinetic energy is energy of motion.")
        self.assertEqual(StudentKeyServer.keys, ["Bearer sk-student"])

    def test_teacher_server_never_forwards_to_itself(self):
        make_workspace(self.root, 'ai = "teacher"\nserver = "http://127.0.0.1:9"')
        run_cli("build", "-w", str(self.root))
        os.environ.update({"LLM_PROVIDER": "openai", "OPENAI_BASE_URL": self.base_url, "OPENAI_API_KEY": "sk-teacher"})
        workspace = workspace_module.load(self.root)
        client = create_explorer(workspace).test_client()
        StudentKeyServer.keys.clear()
        events = stream_events(client, {"query": "What is kinetic energy?"})
        self.assertEqual(events[-1]["type"], "done")
        self.assertEqual(StudentKeyServer.keys, ["Bearer sk-teacher"])

    def test_widget_accepts_only_course_and_local_origins(self):
        make_workspace(self.root, 'ai = "teacher"\nallowed_origins = ["https://lms.example.edu"]')
        config = self.root / "wendao.toml"
        config.write_text(config.read_text(encoding="utf-8").replace('website = ""', 'website = "https://course.example/book/"'), encoding="utf-8")
        run_cli("build", "-w", str(self.root))
        client = create_widget(workspace_module.load(self.root)).test_client()
        for origin, allowed in [("https://course.example", True), ("https://lms.example.edu", True),
                                ("http://127.0.0.1:8000", True), ("https://evil.example", False)]:
            response = client.options("/api/answer/stream", headers={"Origin": origin})
            self.assertEqual(response.headers.get("Access-Control-Allow-Origin") == origin, allowed, origin)

    def test_damaged_or_unsafe_files_are_rejected(self):
        bad = Path(self.tmp.name) / "bad.wendao"
        bad.write_text("not a zip", encoding="utf-8")
        with self.assertRaisesRegex(PackError, "not a Wendao course file"):
            unpack(bad)
        evil = Path(self.tmp.name) / "evil.wendao"
        with zipfile.ZipFile(evil, "w") as archive:
            archive.writestr("course.json", "{}")
            archive.writestr("../escape.txt", "x")
        with self.assertRaisesRegex(PackError, "unsafe"):
            unpack(evil)

    def test_teacher_commands_explain_missing_tools(self):
        make_workspace(self.root, 'ai = "teacher"')
        real_find_spec = __import__("importlib.util").util.find_spec
        with mock.patch("importlib.util.find_spec", lambda name: None if name == "pypdf" else real_find_spec(name)):
            errors = io.StringIO()
            with contextlib.redirect_stderr(errors), self.assertRaises(SystemExit):
                run_cli("build", "-w", str(self.root))
        self.assertIn('pip install "wendao[teacher]"', errors.getvalue())


if __name__ == "__main__":
    unittest.main()


class SignInTest(unittest.TestCase):
    """Roster sign-in: only enrolled emails, limits counted per student, tokens can't be forged."""

    @classmethod
    def setUpClass(cls):
        cls.server = HTTPServer(("127.0.0.1", 0), StudentKeyServer)
        threading.Thread(target=cls.server.serve_forever, daemon=True).start()
        cls.base_url = f"http://127.0.0.1:{cls.server.server_port}/v1"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "physics"
        self.env = mock.patch.dict(os.environ, {
            "WENDAO_CACHE": str(Path(self.tmp.name) / "cache"),
            "LLM_PROVIDER": "openai", "OPENAI_BASE_URL": self.base_url, "OPENAI_API_KEY": "sk-teacher",
        }, clear=True)
        self.env.start()
        make_workspace(self.root, 'ai = "teacher"\nroster = "students.csv"\nquestions_per_day = 2')
        (self.root / "students.csv").write_text(
            "Email,Name,Limit\nAda@uni.edu,Ada Lovelace,\nbo@uni.edu,Bo,3\n", encoding="utf-8"
        )
        run_cli("build", "-w", str(self.root))
        self.workspace = workspace_module.load(self.root)

    def tearDown(self):
        self.env.stop()
        self.tmp.cleanup()

    def sign_in(self, client, email):
        return client.post("/api/login", json={"email": email})

    def ask(self, client, token=""):
        response = client.post("/api/answer/stream", json={"query": "What is kinetic energy?"},
                               headers={"Authorization": f"Bearer {token}"} if token else {})
        return [json.loads(line) for line in response.get_data(as_text=True).splitlines() if line.strip()][-1]

    def test_only_roster_students_can_use_the_course_ai(self):
        client = create_explorer(self.workspace).test_client()
        self.assertTrue(client.get("/api/health").get_json()["ai"]["login_required"])
        self.assertEqual(self.ask(client)["error"], "LoginRequired")

        refused = self.sign_in(client, "stranger@uni.edu")
        self.assertEqual(refused.status_code, 403)
        self.assertIn("not on the class list", refused.get_json()["message"])

        reply = self.sign_in(client, "  ADA@uni.edu ").get_json()
        self.assertEqual((reply["email"], reply["name"]), ("ada@uni.edu", "Ada Lovelace"))
        self.assertEqual(self.ask(client, reply["token"])["type"], "done")
        self.assertEqual(self.ask(client, reply["token"])["type"], "done")
        third = self.ask(client, reply["token"])
        self.assertEqual(third["error"], "AiUnavailable")
        self.assertIn("all 2 of today's questions", third["message"])

        bo = self.sign_in(client, "bo@uni.edu").get_json()["token"]  # personal limit of 3, counted separately
        self.assertEqual([self.ask(client, bo)["type"] for _ in range(3)], ["done", "done", "done"])
        self.assertEqual(self.ask(client, bo)["error"], "AiUnavailable")

        forged = reply["token"].split(".")[0] + ".0123456789abcdef0123456789abcdef"
        self.assertEqual(self.ask(client, forged)["error"], "LoginRequired")

        output = run_cli("students", "-w", str(self.root))
        self.assertIn("2 students", output)
        self.assertRegex(output, r"ada@uni.edu\s+2/2\s+Ada Lovelace")
        self.assertRegex(output, r"bo@uni.edu\s+3/3")

    def test_removing_a_student_from_the_roster_signs_them_out(self):
        client = create_explorer(self.workspace).test_client()
        token = self.sign_in(client, "bo@uni.edu").get_json()["token"]
        roster = self.root / "students.csv"
        roster.write_text("email\nada@uni.edu\n", encoding="utf-8")
        os.utime(roster, (roster.stat().st_mtime + 5,) * 2)
        self.assertEqual(self.ask(client, token)["error"], "LoginRequired")

    def test_widget_uses_the_same_sign_in(self):
        client = create_widget(self.workspace).test_client()
        self.assertEqual(client.post("/api/answer", json={"query": "What is kinetic energy?"}).status_code, 401)
        token = client.post("/api/login", json={"email": "ada@uni.edu"}).get_json()["token"]
        answered = client.post("/api/answer", json={"query": "What is kinetic energy?"}, headers={"Authorization": f"Bearer {token}"})
        self.assertEqual(answered.status_code, 200)

    def test_course_app_signs_in_through_the_teachers_website(self):
        config = self.root / "wendao.toml"
        config.write_text(config.read_text(encoding="utf-8").replace('ai = "teacher"', 'ai = "teacher"\nserver = "http://127.0.0.1:9"'),
                          encoding="utf-8")
        run_cli("pack", "-w", str(self.root))
        course = open_pack(self.root / "build" / "physics.wendao")
        self.assertTrue(course.server_needs_login)
        app_client = create_explorer(course).test_client()
        self.assertTrue(app_client.get("/api/health").get_json()["ai"]["login_required"])
        teacher_site = create_explorer(workspace_module.load(self.root)).test_client()

        def fake_forward_login(server, email):
            return teacher_site.post("/api/login", json={"email": email}).get_json()

        with mock.patch("wendao.web.explorer.forward_login", fake_forward_login):
            reply = app_client.post("/api/login", json={"email": "ada@uni.edu"})
        self.assertEqual(reply.status_code, 200)
        self.assertTrue(reply.get_json()["token"])

    def test_bad_roster_is_reported(self):
        (self.root / "students.csv").write_text("name\nAda\n", encoding="utf-8")
        errors = io.StringIO()
        with contextlib.redirect_stderr(errors), self.assertRaises(SystemExit):
            run_cli("students", "-w", str(self.root))
        self.assertIn("needs an `email` column", errors.getvalue())


def _read_secret(folder: str) -> bytes:
    from wendao.web.accounts import TokenSigner

    return TokenSigner.for_folder(Path(folder)).secret


class RedirectServer(BaseHTTPRequestHandler):
    def do_POST(self):  # noqa: N802
        self.rfile.read(int(self.headers["Content-Length"]))
        self.send_response(302)
        self.send_header("Location", "http://169.254.169.254/latest/meta-data/")
        self.end_headers()

    def log_message(self, *args):
        pass


class ReviewFixesTest(unittest.TestCase):
    """Fixes from the 0.2.0 code review."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "physics"

    def tearDown(self):
        self.tmp.cleanup()

    def test_servers_refuse_private_ai_addresses(self):
        from wendao.web.ai import check_public_endpoint

        for url in ["http://example.com/v1", "https://127.0.0.1/v1", "https://localhost:11434/v1",
                    "https://169.254.169.254/", "https://10.0.0.5/v1", "https://[::1]/v1"]:
            with self.subTest(url=url), self.assertRaises(AiUnavailable):
                check_public_endpoint(url)
        with mock.patch("socket.getaddrinfo", return_value=[(None, None, None, "", ("93.184.216.34", 443))]):
            check_public_endpoint("https://api.example.com/v1")  # public: allowed

        ollama = {"ai": {"provider": "openai", "base_url": "http://127.0.0.1:11434/v1"}}
        with self.assertRaisesRegex(AiUnavailable, "https://"):
            AiPolicy.from_settings("either").choose(ollama, "s")  # course website: refused
        how, _ = AiPolicy.from_settings("either", use_local_key=False).choose(ollama, "s")  # course app: fine
        self.assertEqual(how, "provider")

    def test_student_endpoints_do_not_follow_redirects(self):
        server = HTTPServer(("127.0.0.1", 0), RedirectServer)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        try:
            from wendao.rag.providers import provider_from_name

            provider = provider_from_name("openai", model="m", api_key="sk", base_url=f"http://127.0.0.1:{server.server_port}/v1")
            with self.assertRaisesRegex(RuntimeError, "redirects aren't allowed"):
                provider.generate({"messages": [{"content": "s"}, {"content": "q"}], "evidence": []})
        finally:
            server.shutdown()
            server.server_close()

    def test_daily_limit_is_shared_by_all_server_workers(self):
        fake = HTTPServer(("127.0.0.1", 0), StudentKeyServer)
        threading.Thread(target=fake.serve_forever, daemon=True).start()
        env = {"LLM_PROVIDER": "openai", "OPENAI_BASE_URL": f"http://127.0.0.1:{fake.server_port}/v1", "OPENAI_API_KEY": "sk"}
        try:
            with mock.patch.dict(os.environ, env, clear=True):
                make_workspace(self.root, 'ai = "teacher"\nquestions_per_day = 2')
                run_cli("build", "-w", str(self.root))
                worker_a = create_explorer(workspace_module.load(self.root)).test_client()
                worker_b = create_explorer(workspace_module.load(self.root)).test_client()
                results = [stream_events(worker, {"query": "What is kinetic energy?"})[-1]
                           for worker in (worker_a, worker_b, worker_a)]
        finally:
            fake.shutdown()
            fake.server_close()
        self.assertEqual([event["type"] for event in results], ["done", "done", "error"])
        self.assertIn("all 2 of today's questions", results[2]["message"])

    def test_workers_starting_together_share_one_secret(self):
        import multiprocessing

        folder = Path(self.tmp.name)
        with mock.patch.dict(os.environ, {}, clear=True):
            with multiprocessing.get_context("spawn").Pool(8) as pool:
                secrets_seen = set(pool.map(_read_secret, [str(folder)] * 8))
        self.assertEqual(len(secrets_seen), 1)
        self.assertEqual(sorted(path.name for path in folder.iterdir()), [".wendao-secret"])

    def test_table_of_contents_respects_file_types(self):
        from wendao.ingest import ContentExtractor

        notes = Path(self.tmp.name) / "notes"
        notes.mkdir()
        (notes / "a.md").write_text("# A\n\nMarkdown page", encoding="utf-8")
        (notes / "b.txt").write_text("Plain text page", encoding="utf-8")
        extractor = ContentExtractor(notes, suffixes={".md"})
        extractor.process_all([{"file": "a.md"}, {"file": "b.txt"}])
        self.assertEqual({chunk["file_path"] for chunk in extractor.chunks}, {"a.md"})

    def test_word_tables_stay_in_place(self):
        from docx import Document

        from wendao.readers import read_docx

        path = Path(self.tmp.name) / "lab.docx"
        document = Document()
        document.add_heading("Safety", 1)
        document.add_paragraph("Before the table.")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Item", "Goggles"
        document.add_paragraph("After the table.")
        document.save(str(path))
        text = read_docx(path).parts[0].text
        self.assertLess(text.index("Before the table."), text.index("Item | Goggles"))
        self.assertLess(text.index("Item | Goggles"), text.index("After the table."))


class PinnedEndpointTest(unittest.TestCase):
    """A checked student AI address is pinned, so DNS can't re-point it at a private machine."""

    PUBLIC = [(None, None, None, "", ("93.184.216.34", 443))]

    def test_policy_pins_the_checked_address(self):
        with mock.patch("socket.getaddrinfo", return_value=self.PUBLIC):
            _, provider = AiPolicy.from_settings("either").choose(
                {"ai": {"provider": "openai", "api_key": "sk", "base_url": "https://api.example.com/v1"}}, "s")
        self.assertEqual(provider.pinned_ip, "93.184.216.34")
        with self.assertRaisesRegex(AiUnavailable, "Gemini"):
            AiPolicy.from_settings("either").choose(
                {"ai": {"provider": "gemini", "api_key": "k", "base_url": "https://example.com"}}, "s")

    def test_openai_connects_to_the_pinned_ip_without_looking_up_the_name(self):
        from wendao.rag import providers

        provider = providers.OpenAICompatibleProvider(
            model="m", api_key="sk", base_url="https://api.example.com/v1", use_env=False,
            follow_redirects=False, pinned_ip="93.184.216.34")
        connected = []

        def fake_connect(address, *args, **kwargs):
            connected.append(address)
            raise OSError("stop here")

        with mock.patch.object(providers.socket, "create_connection", fake_connect), \
                mock.patch("socket.getaddrinfo", side_effect=AssertionError("the name must not be looked up again")):
            with self.assertRaises(RuntimeError):
                provider.generate({"messages": [{"content": "s"}, {"content": "q"}], "evidence": []})
        self.assertEqual(connected, [("93.184.216.34", 443)])

    def test_anthropic_tells_curl_to_use_the_pinned_ip(self):
        from wendao.rag import providers

        provider = providers.AnthropicProvider(
            model="m", api_key="sk", base_url="https://relay.example.com", use_env=False, pinned_ip="93.184.216.34")
        seen = {}

        def fake_run(args, **kwargs):
            seen["config"] = Path(args[args.index("--config") + 1]).read_text(encoding="utf-8")
            raise OSError("stop here")

        with mock.patch.object(providers.subprocess, "run", fake_run), self.assertRaises(OSError):
            provider.generate({"messages": [{"content": "s"}, {"content": "q"}], "evidence": []})
        self.assertIn('resolve = "relay.example.com:443:93.184.216.34"', seen["config"])
