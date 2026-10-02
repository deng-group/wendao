from __future__ import annotations

import contextlib
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from wendao.cli import main
from wendao.readers import latex_to_text, read_docx, read_html, read_latex, read_pdf, read_pptx, read_text
from wendao.web.explorer import course_source_url


def make_pdf(path: Path, pages: list[str]) -> None:
    """Write a minimal text PDF (one line of Helvetica per page) without extra libraries."""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        content_id = len(objects)
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
                       "/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Phase Diagrams"
    slide.placeholders[1].text = "A convex hull shows which phases are stable"
    slide.notes_slide.notes_text_frame.text = "Remind students about formation energy"
    second = deck.slides.add_slide(deck.slide_layouts[5])
    second.shapes.title.text = "Energy Table"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Phase", "Energy"
    table.cell(1, 0).text, table.cell(1, 1).text = "Anatase", "-0.2"
    deck.save(str(path))


def make_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Lab Manual", 0)
    document.add_heading("Safety", 1)
    document.add_paragraph("Always wear goggles in the lab.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Item", "Goggles"
    document.save(str(path))


class ReaderTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_pdf_keeps_page_numbers(self):
        make_pdf(self.dir / "notes.pdf", ["Diffusion follows Fick law", "Second page about entropy"])
        document = read_pdf(self.dir / "notes.pdf")
        self.assertEqual([part.location for part in document.parts], ["page 1", "page 2"])
        self.assertIn("Fick law", document.parts[0].text)
        self.assertEqual(document.file_type, "pdf")

    def test_pptx_reads_titles_bullets_tables_and_notes(self):
        make_pptx(self.dir / "week1.pptx")
        document = read_pptx(self.dir / "week1.pptx")
        self.assertEqual([part.location for part in document.parts], ["slide 1", "slide 2"])
        first, second = document.parts[0].text, document.parts[1].text
        self.assertTrue(first.startswith("## Phase Diagrams"))
        self.assertIn("convex hull", first)
        self.assertIn("Speaker notes: Remind students", first)
        self.assertIn("Anatase | -0.2", second)
        self.assertEqual(document.title, "Phase Diagrams")

    def test_docx_turns_headings_into_markdown(self):
        make_docx(self.dir / "manual.docx")
        document = read_docx(self.dir / "manual.docx")
        text = document.parts[0].text
        self.assertEqual(document.title, "Lab Manual")
        self.assertIn("## Safety", text)
        self.assertIn("Always wear goggles", text)
        self.assertIn("Item | Goggles", text)

    def test_latex_keeps_sections_and_math_and_drops_markup(self):
        title, text = latex_to_text(
            r"""\documentclass{article}\title{Thermodynamics}
            \begin{document}\maketitle
            \section{Free energy} % a comment
            The \textbf{Gibbs} energy is $G = H - TS$.\label{eq:g}
            \begin{itemize}\item first \item second\end{itemize}
            \end{document}"""
        )
        self.assertEqual(title, "Thermodynamics")
        self.assertIn("## Free energy", text)
        self.assertIn("The Gibbs energy is $G = H - TS$.", text)
        self.assertIn("- first", text)
        self.assertNotIn("comment", text)
        self.assertNotIn("label", text)

    def test_html_and_text(self):
        (self.dir / "page.html").write_text(
            "<html><head><title>Bonding</title><script>var x=1;</script></head>"
            "<body><nav>menu</nav><h1>Ionic bonds</h1><p>Electrons transfer.</p><ul><li>NaCl</li></ul></body></html>",
            encoding="utf-8",
        )
        document = read_html(self.dir / "page.html")
        self.assertEqual(document.title, "Bonding")
        text = document.parts[0].text
        self.assertIn("# Ionic bonds", text)
        self.assertIn("- NaCl", text)
        self.assertNotIn("var x", text)
        self.assertNotIn("menu", text)

        (self.dir / "notes.txt").write_text("Plain   notes\n\n\n\nhere", encoding="utf-8")
        self.assertEqual(read_text(self.dir / "notes.txt").parts[0].text, "Plain notes\n\nhere")
        self.assertEqual(read_latex.__name__, "read_latex")

    def test_pdf_source_links_open_the_right_page(self):
        self.assertEqual(
            course_source_url("week1/notes.pdf", "https://course.example/", "page 3"),
            "https://course.example/week1/notes.pdf#page=3",
        )
        self.assertEqual(
            course_source_url("structures/crystal_structure.ipynb", "https://course.example"),
            "https://course.example/structures/crystal-structure/",
        )


class MixedFormatCourseTest(unittest.TestCase):
    """A course made of slides, a PDF, and a Word file builds and searches end to end."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "materials-101"
        self.env = mock.patch.dict(os.environ, {}, clear=True)
        self.env.start()

    def tearDown(self):
        self.env.stop()
        self.tmp.cleanup()

    def run_cli(self, *argv: str) -> str:
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            main(list(argv))
        return output.getvalue()

    def test_build_and_search_mixed_formats(self):
        self.run_cli("init", str(self.root))
        notes = self.root / "notes"
        (notes / "week1").mkdir()
        (notes / "week2").mkdir()
        make_pptx(notes / "week1" / "phases.pptx")
        make_pdf(notes / "week2" / "diffusion.pdf", ["Diffusion follows Fick law", "Entropy always increases"])
        make_docx(notes / "lab.docx")
        (notes / "empty.pdf").write_bytes(b"not really a pdf")
        (self.root / "concepts.json").write_text(json.dumps({"version": 1, "concepts": [
            {"id": "convex-hull", "label": "Convex Hull", "category": "thermo", "aliases": ["convex hull"]},
            {"id": "diffusion", "label": "Diffusion", "category": "kinetics", "aliases": ["diffusion"]},
        ]}), encoding="utf-8")
        config = self.root / "wendao.toml"
        config.write_text(config.read_text(encoding="utf-8").replace("[search]", '[search]\nembedding_model = "tfidf"'), encoding="utf-8")

        output = self.run_cli("build", "-w", str(self.root))
        self.assertIn("(pdf, slides, word)", output)
        self.assertIn("Warning: Could not read empty.pdf", output)

        output = self.run_cli("ask", "-w", str(self.root), "--search-only", "what does a convex hull show")
        self.assertIn("1. week1/phases.pptx, slide 1", output)
        output = self.run_cli("ask", "-w", str(self.root), "--search-only", "Fick law diffusion")
        self.assertIn("1. week2/diffusion.pdf, page 1", output)

        chunks = [json.loads(line) for line in (self.root / "build" / "chunks.jsonl").read_text().splitlines()]
        ids = [chunk["chunk_id"] for chunk in chunks]
        self.assertEqual(len(ids), len(set(ids)), "chunk ids must be unique across pages and slides")

    def test_file_types_setting_limits_what_is_read(self):
        self.run_cli("init", str(self.root))
        make_pdf(self.root / "notes" / "a.pdf", ["Some text"])
        (self.root / "notes" / "b.md").write_text("# B\n\nMarkdown text", encoding="utf-8")
        config = self.root / "wendao.toml"
        config.write_text(config.read_text(encoding="utf-8").replace('# file_types = ["pdf", "pptx"]', 'file_types = ["md"]'),
                          encoding="utf-8")
        output = self.run_cli("extract", "-w", str(self.root))
        self.assertIn("1 files (markdown)", output)

        config.write_text(config.read_text(encoding="utf-8").replace('file_types = ["md"]', 'file_types = ["md", "xlsx"]'),
                          encoding="utf-8")
        errors = io.StringIO()
        with contextlib.redirect_stderr(errors), self.assertRaises(SystemExit):
            self.run_cli("extract", "-w", str(self.root))
        self.assertIn("Unsupported file_types", errors.getvalue())


if __name__ == "__main__":
    unittest.main()
