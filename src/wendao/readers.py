"""Read lecture material in different file formats as plain, Markdown-style text.

Each reader returns a `Document`: a title and a list of parts. A part is a piece of text
plus where it came from in the file (for example "page 12" or "slide 5"), so answers can
point students to the exact page. Headings are written as Markdown (`# Title`) so the
chunker can split on them the same way for every format.

Supported: PDF, PowerPoint (.pptx), Word (.docx), LaTeX (.tex), HTML, and plain text.
Markdown and Jupyter notebooks are handled in `ingest.py`.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path


@dataclass
class Part:
    text: str
    location: str = ""


@dataclass
class Document:
    title: str
    file_type: str
    parts: list[Part] = field(default_factory=list)


def clean(text: str) -> str:
    """Normalize whitespace while keeping paragraph breaks."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# PDF ------------------------------------------------------------------------------------------


def read_pdf(path: Path) -> Document:
    import logging

    from pypdf import PdfReader

    logging.getLogger("pypdf").setLevel(logging.ERROR)  # problems are reported as build warnings instead

    reader = PdfReader(str(path))
    title = ""
    try:
        title = (reader.metadata.title or "").strip() if reader.metadata else ""
    except Exception:  # noqa: BLE001 - broken metadata should not stop extraction
        title = ""
    parts = []
    for number, page in enumerate(reader.pages, start=1):
        text = clean(page.extract_text() or "")
        if text:
            parts.append(Part(text, f"page {number}"))
    if not title and parts:
        first_line = parts[0].text.splitlines()[0].strip()
        title = first_line if 3 <= len(first_line) <= 120 else ""
    return Document(title or path.stem, "pdf", parts)


# PowerPoint -----------------------------------------------------------------------------------


def _shape_text(shape) -> list[str]:
    lines = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if text:
                lines.append(("  " * paragraph.level) + "- " + text if paragraph.level else text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    for child in getattr(shape, "shapes", []) or []:  # grouped shapes
        lines.extend(_shape_text(child))
    return lines


def read_pptx(path: Path) -> Document:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts = []
    deck_title = ""
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        slide_title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        lines = []
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            lines.extend(_shape_text(shape))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        body = "\n".join(lines)
        if notes:
            body = f"{body}\n\nSpeaker notes: {notes}" if body else f"Speaker notes: {notes}"
        heading = f"## {slide_title}" if slide_title else f"## Slide {number}"
        text = clean(f"{heading}\n\n{body}")
        if body.strip() or slide_title:
            parts.append(Part(text, f"slide {number}"))
        if not deck_title and slide_title:
            deck_title = slide_title
    core_title = (presentation.core_properties.title or "").strip()
    return Document(core_title or deck_title or path.stem, "slides", parts)


# Word -----------------------------------------------------------------------------------------


def read_docx(path: Path) -> Document:
    from docx import Document as WordDocument

    document = WordDocument(str(path))
    blocks = []
    title = (document.core_properties.title or "").strip()
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        match = re.match(r"heading (\d)", style)
        if style == "title":
            title = title or text
            blocks.append(f"# {text}")
        elif match:
            blocks.append("#" * min(int(match.group(1)) + 1, 6) + f" {text}")
        else:
            blocks.append(text)
    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        rows = [row for row in rows if row.replace("|", "").strip()]
        if rows:
            blocks.append("\n".join(rows))
    if not title:
        first_heading = next((block.lstrip("# ") for block in blocks if block.startswith("#")), "")
        title = first_heading
    return Document(title or path.stem, "word", [Part(clean("\n\n".join(blocks)))] if blocks else [])


# LaTeX ----------------------------------------------------------------------------------------

LATEX_HEADINGS = {"part": "#", "chapter": "#", "section": "##", "subsection": "###", "subsubsection": "####"}


def latex_to_text(source: str) -> tuple[str, str]:
    """Turn LaTeX into readable Markdown-style text. Returns (title, text). Math is kept as written."""
    source = re.sub(r"(?<!\\)%.*", "", source)  # comments
    body = source.split(r"\begin{document}", 1)[1] if r"\begin{document}" in source else source
    body = body.split(r"\end{document}", 1)[0]
    title_match = re.search(r"\\title\{([^{}]*)\}", source)
    title = title_match.group(1).strip() if title_match else ""

    def heading(match: re.Match) -> str:
        return f"\n\n{LATEX_HEADINGS[match.group(1)]} {match.group(2).strip()}\n\n"

    body = re.sub(r"\\(part|chapter|section|subsection|subsubsection)\*?\{([^{}]*)\}", heading, body)
    body = re.sub(r"\\item\s*", "\n- ", body)
    body = re.sub(r"\\(label|ref|eqref|cite|citep|citet|pageref|includegraphics|vspace|hspace)(\[[^\]]*\])?\{[^{}]*\}", "", body)
    body = re.sub(r"\\(begin|end)\{(itemize|enumerate|description|center|flushleft|flushright|figure|table|minipage)\*?\}(\[[^\]]*\])?", "", body)
    for _ in range(3):  # unwrap simple formatting commands, innermost first
        body = re.sub(r"\\(textbf|textit|emph|underline|texttt|textrm|textsf|caption|footnote|url|mbox)\{([^{}]*)\}", r"\2", body)
    body = re.sub(r"\\(maketitle|tableofcontents|newpage|clearpage|centering|noindent|small|large|Large|footnotesize)\b", "", body)
    body = body.replace("~", " ").replace(r"\\", "\n")
    return title, clean(body)


def read_latex(path: Path) -> Document:
    title, text = latex_to_text(path.read_text(encoding="utf-8", errors="replace"))
    return Document(title or path.stem, "latex", [Part(text)] if text else [])


# HTML -----------------------------------------------------------------------------------------


class _HtmlText(HTMLParser):
    SKIP = {"script", "style", "nav", "header", "footer", "noscript", "svg"}
    BLOCK = {"p", "div", "section", "article", "li", "tr", "br", "table", "ul", "ol", "pre", "blockquote"}

    def __init__(self):
        super().__init__()
        self.out: list[str] = []
        self.skip_depth = 0
        self.title = ""
        self.in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip_depth += 1
        elif tag == "title":
            self.in_title = True
        elif re.fullmatch(r"h[1-6]", tag):
            self.out.append("\n\n" + "#" * int(tag[1]) + " ")
        elif tag == "li":
            self.out.append("\n- ")
        elif tag in self.BLOCK:
            self.out.append("\n\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip_depth:
            self.skip_depth -= 1
        elif tag == "title":
            self.in_title = False
        elif re.fullmatch(r"h[1-6]", tag) or tag in self.BLOCK:
            self.out.append("\n\n")

    def handle_data(self, data):
        if self.in_title:
            self.title += data
        elif not self.skip_depth:
            self.out.append(data)


def read_html(path: Path) -> Document:
    parser = _HtmlText()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    text = clean(re.sub(r"[ \t]*\n[ \t]*\n[ \t]*", "\n\n", "".join(parser.out)))
    return Document(parser.title.strip() or path.stem, "html", [Part(text)] if text else [])


# Plain text -----------------------------------------------------------------------------------


def read_text(path: Path) -> Document:
    text = clean(path.read_text(encoding="utf-8", errors="replace"))
    return Document(path.stem, "text", [Part(text)] if text else [])


READERS = {
    ".pdf": read_pdf,
    ".pptx": read_pptx,
    ".docx": read_docx,
    ".tex": read_latex,
    ".html": read_html,
    ".htm": read_html,
    ".txt": read_text,
}
